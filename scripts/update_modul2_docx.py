"""
Script to update D:/JoyOps/docs/Modul2/Test Design Specification & TTM - JoyOps.docx
Adds comprehensive positive (happy path) and negative (invalid/boundary/error) test conditions,
coverage items, traceability matrix (TTM), and test design technique summaries.
"""

import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

def set_cell_properties(cell, fill_hex=None, width_dxa=None, bold=False, font_size_pt=9.5, align="left"):
    tcPr = cell._tc.get_or_add_tcPr()
    
    # Width
    if width_dxa:
        tcW = tcPr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW')
            tcPr.append(tcW)
        tcW.set(qn('w:w'), str(width_dxa))
        tcW.set(qn('w:type'), 'dxa')
    
    # Borders
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
    
    # Shading (background)
    shd = tcPr.find(qn('w:shd'))
    if fill_hex:
        if shd is None:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
            tcPr.append(shd)
        else:
            shd.set(qn('w:fill'), fill_hex)
    elif shd is not None:
        tcPr.remove(shd)
        
    # Margins
    tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="80" w:type="dxa"/>'
            f'<w:left w:w="120" w:type="dxa"/>'
            f'<w:bottom w:w="80" w:type="dxa"/>'
            f'<w:right w:w="120" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)
        
    # Format paragraph & run inside cell
    for p in cell.paragraphs:
        if align == "center":
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(font_size_pt)
            r.bold = bold

def populate_table(table, headers, rows_data, col_widths, center_cols=set()):
    # Clear existing rows except the first one
    while len(table.rows) > 1:
        tr = table.rows[-1]._tr
        table._tbl.remove(tr)
        
    # Set header (Row 0)
    hdr_row = table.rows[0]
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.text = title
        align = "center" if c_idx in center_cols else "left"
        set_cell_properties(cell, fill_hex="A4C2F4", width_dxa=col_widths[c_idx], bold=True, font_size_pt=9.5, align=align)
        
    # Add data rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.add_row()
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(val)
            align = "center" if c_idx in center_cols else "left"
            set_cell_properties(cell, fill_hex=None, width_dxa=col_widths[c_idx], bold=False, font_size_pt=9.0, align=align)

def main():
    docx_path = 'D:/JoyOps/docs/Modul2/Test Design Specification & TTM - JoyOps.docx'
    doc = docx.Document(docx_path)
    
    # --- TABLE 0: TEST CONDITIONS ---
    t0_headers = ['ID', 'Feature Set', 'Description', 'Priority', 'Traceability']
    t0_widths = [1100, 1100, 4700, 1000, 1100]
    t0_center = {0, 1, 3, 4}
    t0_data = [
        # FS-001 (RQ-001)
        ['TCond-001', 'FS-001', 'Login dengan username dan kata sandi yang valid pada akun Admin yang sudah aktif dan disetujui developer.', 'High', 'RQ-001'],
        ['TCond-002', 'FS-001', 'Login dengan kombinasi kata sandi salah atau username yang tidak terdaftar di sistem.', 'High', 'RQ-001'],
        ['TCond-003', 'FS-001', 'Login menggunakan akun Admin yang berstatus Menunggu (belum disetujui) atau Ditolak oleh developer.', 'High', 'RQ-001'],
        ['TCond-004', 'FS-001', 'Login dengan mengosongkan salah satu atau kedua kolom isian wajib (username atau kata sandi kosong).', 'Medium', 'RQ-001'],
        # FS-002 (RQ-002)
        ['TCond-005', 'FS-002', 'Pendaftaran akun Admin baru dengan username unik, nama lengkap, dan kata sandi valid (memenuhi kriteria minimum).', 'High', 'RQ-002'],
        ['TCond-006', 'FS-002', 'Pendaftaran akun Admin baru dengan username yang sudah terdaftar atau sedang dalam status Menunggu persetujuan.', 'High', 'RQ-002'],
        ['TCond-007', 'FS-002', 'Pendaftaran akun Admin baru dengan panjang kata sandi kurang dari 6 karakter atau field wajib yang dikosongkan.', 'Medium', 'RQ-002'],
        # FS-003 (RQ-003)
        ['TCond-008', 'FS-003', 'Pengajuan permohonan reset sandi untuk username Admin yang valid dan terdaftar di sistem.', 'High', 'RQ-003'],
        ['TCond-009', 'FS-003', 'Pengajuan permohonan reset sandi menggunakan username yang tidak terdaftar di sistem.', 'High', 'RQ-003'],
        ['TCond-010', 'FS-003', 'Mengakses form reset sandi menggunakan tautan/token penggantian sandi yang sudah kedaluwarsa atau pernah digunakan.', 'Medium', 'RQ-003'],
        # FS-004 (RQ-004)
        ['TCond-011', 'FS-004', 'Menyimpan paket layanan laundry baru beserta tarif, satuan (kg/pcs), dan durasi pengerjaan yang valid.', 'High', 'RQ-004'],
        ['TCond-012', 'FS-004', 'Menyimpan paket layanan dengan nilai tarif di luar batas yang diizinkan (tarif Rp 0, bernilai negatif, atau > Rp 500.000).', 'High', 'RQ-004'],
        ['TCond-013', 'FS-004', 'Menyimpan paket layanan dengan nilai durasi pengerjaan tidak valid (durasi 0 menit atau kurang dari batas minimal 15 menit).', 'Medium', 'RQ-004'],
        # FS-005 (RQ-005)
        ['TCond-014', 'FS-005', 'Membuat order kasir baru dengan sepasang unit mesin cuci dan mesin pengering berbeda yang berstatus Tersedia serta beban cucian valid.', 'High', 'RQ-005'],
        ['TCond-015', 'FS-005', 'Membuat order baru dengan memilih unit mesin cuci atau mesin pengering yang sedang berstatus tidak tersedia (Digunakan atau Maintenance).', 'High', 'RQ-005'],
        ['TCond-016', 'FS-005', 'Membuat order baru dengan memasukkan beban cucian melebihi batas kapasitas maksimal unit mesin terpilih (Qty > Kapasitas).', 'High', 'RQ-005'],
        ['TCond-017', 'FS-005', 'Membuat order baru dengan nilai beban cucian tidak valid (Qty bernilai 0 atau angka desimal negatif).', 'Medium', 'RQ-005'],
        # FS-006 (RQ-006)
        ['TCond-018', 'FS-006', 'Mengubah status operasional unit mesin antara Tersedia dan Maintenance serta menyimpan kapasitas valid (5–30 kg).', 'Medium', 'RQ-006'],
        ['TCond-019', 'FS-006', 'Menambahkan atau memperbarui unit mesin dengan nilai kapasitas di luar rentang batas operasional (< 5 kg atau > 30 kg).', 'Medium', 'RQ-006'],
        ['TCond-020', 'FS-006', 'Mendaftarkan unit mesin baru dengan kode mesin duplikat atau format kode tidak sesuai standar (bukan awalan WM-*/DM-*).', 'Medium', 'RQ-006'],
        # FS-007 (RQ-007)
        ['TCond-021', 'FS-007', 'Menandai order berstatus Berjalan menjadi Selesai dan melepaskan kedua unit mesin kembali menjadi Tersedia secara otomatis.', 'High', 'RQ-007'],
        ['TCond-022', 'FS-007', 'Mencoba menandai selesai pada order yang sudah berstatus Selesai atau Dibatalkan (transisi status tidak valid).', 'Medium', 'RQ-007'],
        ['TCond-023', 'FS-007', 'Memantau sisa durasi hitung mundur saat timer proses cucian belum berakhir untuk memastikan status tetap terkunci Digunakan.', 'Medium', 'RQ-007'],
        # FS-008 (RQ-008)
        ['TCond-024', 'FS-008', 'Melakukan pelunasan order kasir dengan nominal pembayaran tunai pas atau lebih, menerbitkan struk, dan mencatat mutasi kas masuk.', 'High', 'RQ-008'],
        ['TCond-025', 'FS-008', 'Melakukan pelunasan order dengan nominal pembayaran tunai kurang dari total tagihan cucian (pembayaran tidak mencukupi).', 'High', 'RQ-008'],
        ['TCond-026', 'FS-008', 'Mencoba membatalkan transaksi keuangan atau menghapus rekaman kas oleh pengguna non-Admin tanpa izin hak akses.', 'Medium', 'RQ-008'],
    ]
    populate_table(doc.tables[0], t0_headers, t0_data, t0_widths, t0_center)
    print("Table 0 updated.")

    # --- TABLE 1: COVERAGE ITEMS ---
    t1_headers = ['ID', 'Description', 'Related Test Condition', 'Priority', 'Traceability']
    t1_widths = [1000, 4800, 1500, 1000, 1100]
    t1_center = {0, 2, 3, 4}
    t1_data = [
        ['CI-001', 'Form login: autentikasi kredensial username dan kata sandi valid serta penerbitan sesi Admin', 'TCond-001', 'High', 'RQ-001'],
        ['CI-002', 'Form login: penanganan galat kredensial salah dan username tidak terdaftar', 'TCond-002', 'High', 'RQ-001'],
        ['CI-003', 'Form login: pemblokiran akses akun berstatus Menunggu atau Ditolak developer', 'TCond-003', 'High', 'RQ-001'],
        ['CI-004', 'Form login: validasi pesan isian wajib kosong pada field username dan kata sandi', 'TCond-004', 'Medium', 'RQ-001'],
        ['CI-005', 'Form registrasi Admin: penerbitan permohonan akun berstatus Menunggu', 'TCond-005', 'High', 'RQ-002'],
        ['CI-006', 'Form registrasi Admin: validasi pencegahan duplikasi username terdaftar atau pending', 'TCond-006', 'High', 'RQ-002'],
        ['CI-007', 'Form registrasi Admin: validasi panjang minimum kata sandi (minimal 6 karakter)', 'TCond-007', 'Medium', 'RQ-002'],
        ['CI-008', 'Form lupa sandi: verifikasi username terdaftar dan pencatatan permohonan lupa sandi', 'TCond-008', 'High', 'RQ-003'],
        ['CI-009', 'Form lupa sandi: penanganan pesan galat username tidak ditemukan di sistem', 'TCond-009', 'High', 'RQ-003'],
        ['CI-010', 'Form reset sandi: validasi masa kedaluwarsa dan sifat sekali pakai pada tautan reset', 'TCond-010', 'Medium', 'RQ-003'],
        ['CI-011', 'Form katalog layanan: penyimpanan nama paket, satuan (kg/pcs), tarif, dan durasi valid', 'TCond-011', 'High', 'RQ-004'],
        ['CI-012', 'Form katalog layanan: validasi batas tarif (minimal Rp 1.000 s.d. maksimal Rp 500.000)', 'TCond-012', 'High', 'RQ-004'],
        ['CI-013', 'Form katalog layanan: validasi batas minimal durasi pengerjaan cucian (15 menit)', 'TCond-013', 'Medium', 'RQ-004'],
        ['CI-014', 'Form order kasir: alokasi pasangan unit mesin cuci & pengering berstatus Tersedia', 'TCond-014', 'High', 'RQ-005'],
        ['CI-015', 'Form order kasir: pencegahan pemilihan mesin berstatus Digunakan atau Maintenance', 'TCond-015', 'High', 'RQ-005'],
        ['CI-016', 'Form order kasir: validasi batas kapasitas muatan maksimal unit mesin terpilih', 'TCond-016', 'High', 'RQ-005'],
        ['CI-017', 'Form order kasir: validasi nilai beban cucian positif non-nol (Qty > 0)', 'TCond-017', 'Medium', 'RQ-005'],
        ['CI-018', 'Modul mesin: pembaruan status operasional mesin (Tersedia/Maintenance) dan kapasitas', 'TCond-018', 'Medium', 'RQ-006'],
        ['CI-019', 'Modul mesin: validasi rentang batas kapasitas unit mesin (5 kg s.d. 30 kg)', 'TCond-019', 'Medium', 'RQ-006'],
        ['CI-020', 'Modul mesin: pencegahan duplikasi kode unit mesin dan validasi awalan WM/DM', 'TCond-020', 'Medium', 'RQ-006'],
        ['CI-021', 'Modul order: aksi selesaikan order dan pelepasan otomatis status kedua mesin', 'TCond-021', 'High', 'RQ-007'],
        ['CI-022', 'Modul order: validasi konsistensi status order mencegah eksekusi penyelesaian ganda', 'TCond-022', 'Medium', 'RQ-007'],
        ['CI-023', 'Dashboard kasir: pemantauan sisa waktu timer aktif dan penguncian status mesin', 'TCond-023', 'Medium', 'RQ-007'],
        ['CI-024', 'Modul kasir: perhitungan pelunasan tunai, kembalian tagihan, dan penerbitan struk', 'TCond-024', 'High', 'RQ-008'],
        ['CI-025', 'Modul kasir: validasi nominal pembayaran tunai minimum mencukupi total tagihan', 'TCond-025', 'High', 'RQ-008'],
        ['CI-026', 'Modul keuangan: pembatasan hak akses pembatalan transaksi kas khusus peran Admin', 'TCond-026', 'Medium', 'RQ-008'],
    ]
    populate_table(doc.tables[1], t1_headers, t1_data, t1_widths, t1_center)
    print("Table 1 updated.")

    # --- TABLE 2: TEST TRACEABILITY MATRIX (TTM) ---
    t2_headers = ['Requirement', 'Feature Set', 'Test Condition', 'Coverage Item']
    t2_widths = [1400, 1400, 1800, 4800]
    t2_center = {0, 1, 2}
    t2_data = [
        ['RQ-001', 'FS-001', 'TCond-001', 'CI-001'],
        ['RQ-001', 'FS-001', 'TCond-002', 'CI-002'],
        ['RQ-001', 'FS-001', 'TCond-003', 'CI-003'],
        ['RQ-001', 'FS-001', 'TCond-004', 'CI-004'],
        ['RQ-002', 'FS-002', 'TCond-005', 'CI-005'],
        ['RQ-002', 'FS-002', 'TCond-006', 'CI-006'],
        ['RQ-002', 'FS-002', 'TCond-007', 'CI-007'],
        ['RQ-003', 'FS-003', 'TCond-008', 'CI-008'],
        ['RQ-003', 'FS-003', 'TCond-009', 'CI-009'],
        ['RQ-003', 'FS-003', 'TCond-010', 'CI-010'],
        ['RQ-004', 'FS-004', 'TCond-011', 'CI-011'],
        ['RQ-004', 'FS-004', 'TCond-012', 'CI-012'],
        ['RQ-004', 'FS-004', 'TCond-013', 'CI-013'],
        ['RQ-005', 'FS-005', 'TCond-014', 'CI-014'],
        ['RQ-005', 'FS-005', 'TCond-015', 'CI-015'],
        ['RQ-005', 'FS-005', 'TCond-016', 'CI-016'],
        ['RQ-005', 'FS-005', 'TCond-017', 'CI-017'],
        ['RQ-006', 'FS-006', 'TCond-018', 'CI-018'],
        ['RQ-006', 'FS-006', 'TCond-019', 'CI-019'],
        ['RQ-006', 'FS-006', 'TCond-020', 'CI-020'],
        ['RQ-007', 'FS-007', 'TCond-021', 'CI-021'],
        ['RQ-007', 'FS-007', 'TCond-022', 'CI-022'],
        ['RQ-007', 'FS-007', 'TCond-023', 'CI-023'],
        ['RQ-008', 'FS-008', 'TCond-024', 'CI-024'],
        ['RQ-008', 'FS-008', 'TCond-025', 'CI-025'],
        ['RQ-008', 'FS-008', 'TCond-026', 'CI-026'],
    ]
    populate_table(doc.tables[2], t2_headers, t2_data, t2_widths, t2_center)
    print("Table 2 updated.")

    # --- TABLE 3: TEST DESIGN TECHNIQUE SUMMARY ---
    t3_headers = ['Test Condition ID', 'Test Design Technique', 'Rationale']
    t3_widths = [1600, 2600, 5000]
    t3_center = {0}
    t3_data = [
        ['TCond-001', 'Equivalence Partitioning', 'Menguji partisi data valid untuk kredensial Admin yang sah dan akun berstatus aktif.'],
        ['TCond-002', 'Equivalence Partitioning', 'Menguji partisi data tidak valid untuk username yang tidak terdaftar dan kata sandi yang salah.'],
        ['TCond-003', 'Decision Table Testing', 'Menguji kombinasi kondisi status akun pengguna (disetujui vs pending/ditolak developer) terhadap hak akses login.'],
        ['TCond-004', 'Equivalence Partitioning', 'Menguji partisi masukan kosong (null/empty values) pada field input wajib form login.'],
        ['TCond-005', 'Equivalence Partitioning', 'Menguji partisi data valid pada seluruh field pendaftaran akun Admin baru.'],
        ['TCond-006', 'Equivalence Partitioning', 'Memisahkan partisi username unik dari username yang sudah terdaftar atau dalam antrean persetujuan.'],
        ['TCond-007', 'Boundary Value Analysis', 'Menguji batas minimum panjang karakter kata sandi tepat di batas bawah (6 karakter) dan di bawah batas (< 6 karakter).'],
        ['TCond-008', 'Equivalence Partitioning', 'Menguji masukan username valid yang terdaftar dalam basis data pengguna.'],
        ['TCond-009', 'Equivalence Partitioning', 'Menguji penanganan sistem terhadap partisi masukan username fiktif/tidak terdaftar.'],
        ['TCond-010', 'State Transition Testing', 'Menguji validitas status token reset sandi berdasarkan siklus waktu (aktif vs expired) dan frekuensi penggunaan (unused vs used).'],
        ['TCond-011', 'Equivalence Partitioning', 'Menguji masukan data katalog layanan yang berada dalam rentang spesifikasi operasional.'],
        ['TCond-012', 'Boundary Value Analysis', 'Menguji nilai batas tarif layanan pada ambang batas bawah (Rp 1.000, Rp 999, Rp 0) dan batas atas (Rp 500.000, Rp 500.001).'],
        ['TCond-013', 'Boundary Value Analysis', 'Menguji batas durasi minimal pengerjaan cucian tepat pada ambang bawah (15 menit) dan nilai tidak valid (< 15 menit).'],
        ['TCond-014', 'Decision Table Testing', 'Menguji kombinasi kebenaran status ketersediaan mesin cuci, mesin pengering, dan beban cucian yang valid.'],
        ['TCond-015', 'Decision Table Testing', 'Menguji penolakan alokasi mesin ketika salah satu atau kedua unit mesin bernilai salah/tidak tersedia (Digunakan/Maintenance).'],
        ['TCond-016', 'Boundary Value Analysis', 'Menguji batas kapasitas maksimal muatan mesin pada nilai tepat kapasitas (Qty = Kapasitas) dan melampaui kapasitas (Qty > Kapasitas).'],
        ['TCond-017', 'Boundary Value Analysis', 'Menguji batas bawah nilai berat cucian pada angka nol (Qty = 0) dan angka desimal negatif (Qty < 0).'],
        ['TCond-018', 'State Transition Testing', 'Menguji transisi status mesin yang sah antara kondisi Tersedia dan Maintenance.'],
        ['TCond-019', 'Boundary Value Analysis', 'Menguji nilai ambang batas kapasitas mesin pada batas bawah (5 kg, 4 kg) dan batas atas (30 kg, 31 kg).'],
        ['TCond-020', 'Equivalence Partitioning', 'Memisahkan partisi kode mesin unik dari kode mesin yang duplikat atau format tidak sesuai.'],
        ['TCond-021', 'State Transition Testing', 'Menguji perpindahan siklus hidup order dari status Berjalan ke Selesai serta transisi status mesin menjadi Tersedia.'],
        ['TCond-022', 'State Transition Testing', 'Menguji pencegahan transisi status yang tidak valid dari order yang sudah berstatus terminal (Selesai/Dibatalkan).'],
        ['TCond-023', 'Boundary Value Analysis', 'Menguji sisa durasi timer cucian pada interval waktu positif sebelum habis untuk memastikan status terkunci.'],
        ['TCond-024', 'Boundary Value Analysis', 'Menguji nilai pembayaran tunai pada batas tepat sama dengan total tagihan dan nilai pembayaran lebih (dengan kembalian).'],
        ['TCond-025', 'Boundary Value Analysis', 'Menguji nilai pembayaran tunai di bawah total tagihan (underpayment) untuk memastikan sistem menolak pelunasan.'],
        ['TCond-026', 'Decision Table Testing', 'Menguji aturan hak akses otorisasi pembatalan transaksi berdasarkan peran pengguna (Admin vs Non-Admin/Karyawan).'],
    ]
    populate_table(doc.tables[3], t3_headers, t3_data, t3_widths, t3_center)
    print("Table 3 updated.")

    doc.save(docx_path)
    print("Document saved successfully to:", docx_path)

if __name__ == '__main__':
    main()
